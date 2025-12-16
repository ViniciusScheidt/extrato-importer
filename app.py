import os
import re
import datetime
import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter import ttk

import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import numbers


# =====================================================
# Utils
# =====================================================

def brl_to_float(s: str) -> float:
    s = s.strip().replace("R$", "").replace(" ", "")
    neg = s.startswith("-")
    s = s.replace("-", "")
    s = s.replace(".", "").replace(",", ".")
    v = float(s)
    return -v if neg else v


def format_cpf(num: str) -> str:
    return f"{num[0:3]}.{num[3:6]}.{num[6:9]}-{num[9:11]}"


def format_cnpj(num: str) -> str:
    return f"{num[0:2]}.{num[2:5]}.{num[5:8]}/{num[8:12]}-{num[12:14]}"


def extract_cpf_or_cnpj(text: str) -> str:
    """
    Regra:
    - Se houver CNPJ (14 dígitos), usa CNPJ
    - Senão, se houver CPF (11 dígitos), usa CPF
    - Senão, retorna vazio
    """
    digits = re.sub(r"\D+", "", text)

    cnpjs = re.findall(r"\d{14}", digits)
    if cnpjs:
        return format_cnpj(cnpjs[0])

    cpfs = re.findall(r"\d{11}", digits)
    if cpfs:
        return format_cpf(cpfs[0])

    return ""


def read_pdf_lines(pdf_path: str) -> list[str]:
    lines: list[str] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for ln in text.splitlines():
                ln = ln.strip()
                if ln:
                    lines.append(ln)
    return lines


def detect_bank(pdf_path: str) -> str:
    blob = "\n".join(read_pdf_lines(pdf_path)[:250]).upper()
    if "SICREDI" in blob:
        return "SICREDI"
    if "PAGBANK" in blob or "PAGSEGURO" in blob:
        return "PAGBANK"
    if "STONE" in blob:
        return "STONE"
    return "UNKNOWN"


# =====================================================
# Parsers
# =====================================================

def parse_sicredi(pdf_path: str, only_outgoing: bool = True) -> list[dict]:
    lines = read_pdf_lines(pdf_path)
    txs: list[dict] = []

    date_re = re.compile(r"^(\d{2}/\d{2}/\d{4})\s+(.*)$")

    for ln in lines:
        m = date_re.match(ln)
        if not m:
            continue

        date_str, rest = m.group(1), m.group(2)
        full = ln.upper()

        if (
            "PAGAMENTO PIX" not in full
            and "LIQUIDACAO BOLETO" not in full
            and "DEBITO CONVENIO" not in full
        ):
            continue

        nums = re.findall(r"-?\d{1,3}(?:\.\d{3})*,\d{2}", ln)
        if not nums:
            continue

        valor_str = nums[-2] if len(nums) >= 2 else nums[-1]
        valor = brl_to_float(valor_str)

        if only_outgoing and valor >= 0:
            continue

        cpf_cnpj = extract_cpf_or_cnpj(rest)

        desc = re.sub(
            r"\s+-?\d{1,3}(?:\.\d{3})*,\d{2}(\s+-?\d{1,3}(?:\.\d{3})*,\d{2})?$",
            "",
            rest,
        )

        dt = datetime.datetime.strptime(date_str, "%d/%m/%Y").date()
        txs.append({
            "date": dt,
            "cpf_cnpj": cpf_cnpj,
            "desc": desc,
            "valor_pago": abs(valor),
        })

    return txs


def parse_pagbank(pdf_path: str, only_outgoing: bool = True, ignore_balance: bool = True) -> list[dict]:
    lines = read_pdf_lines(pdf_path)
    txs: list[dict] = []

    line_re = re.compile(r"^(\d{2}/\d{2}/\d{4})\s+(.*?)\s+R\$\s*([0-9\.\,]+)$")

    for ln in lines:
        m = line_re.match(ln)
        if not m:
            continue

        date_str, desc, val_str = m.group(1), m.group(2), m.group(3)
        if ignore_balance and ("SALDO DO DIA" in desc.upper() or "RENDIMENTO" in desc.upper()):
            continue

        valor = brl_to_float(val_str)
        if only_outgoing and valor >= 0:
            continue

        dt = datetime.datetime.strptime(date_str, "%d/%m/%Y").date()
        txs.append({
            "date": dt,
            "cpf_cnpj": extract_cpf_or_cnpj(desc),
            "desc": desc,
            "valor_pago": abs(valor),
        })

    return txs


def parse_stone(pdf_path: str, only_outgoing: bool = True, ignore_fees: bool = True) -> list[dict]:
    lines = read_pdf_lines(pdf_path)
    txs: list[dict] = []

    row_re = re.compile(
        r"^(\d{2}/\d{2}/\d{2})\s+(Entrada|Saída)\s+(.*?)\s+(-?\s*R\$\s*[0-9\.\,]+)"
    )

    buf: list[str] = []
    for ln in lines:
        buf.append(ln)
        joined = " ".join(buf)

        m = row_re.match(joined)
        if not m:
            if len(buf) > 6:
                buf = buf[-3:]
            continue

        buf = []

        date_str, typ, desc, val_str = m.group(1), m.group(2), m.group(3), m.group(4)
        if ignore_fees and "TARIFA" in desc.upper():
            continue

        valor = brl_to_float(val_str)
        if only_outgoing and typ != "Saída":
            continue

        dt = datetime.datetime.strptime(date_str, "%d/%m/%y").date()
        txs.append({
            "date": dt,
            "cpf_cnpj": extract_cpf_or_cnpj(desc),
            "desc": desc,
            "valor_pago": abs(valor),
        })

    return txs


# =====================================================
# Writer
# =====================================================

REQUIRED_HEADERS = [
    "Número da Nota",
    "CPF/CNPJ do Fornecedor",
    "Data do Extrato Bancário",
    "Data da Baixa da Parcela",
    "Valor Pago",
    "Valor Juros",
    "Valor Multa",
    "Valor Desconto",
    "Código da Conta Banco/Caixa",
    "Caixa/Banco",
]


def fill_sheet(xlsx_path, out_path, sheet_name, txs, bank_code, caixa_banco):
    wb = load_workbook(xlsx_path)
    ws = wb[sheet_name]

    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers) if h}

    for r in range(2, ws.max_row + 1):
        for c in range(1, 11):
            ws.cell(row=r, column=c).value = None

    row = 2
    for tx in txs:
        ws.cell(row=row, column=col["CPF/CNPJ do Fornecedor"]).value = tx["cpf_cnpj"]
        ws.cell(row=row, column=col["Data do Extrato Bancário"]).value = tx["date"]
        ws.cell(row=row, column=col["Data da Baixa da Parcela"]).value = tx["date"]
        ws.cell(row=row, column=col["Valor Pago"]).value = tx["valor_pago"]
        ws.cell(row=row, column=col["Código da Conta Banco/Caixa"]).value = bank_code
        ws.cell(row=row, column=col["Caixa/Banco"]).value = caixa_banco
        row += 1

    wb.save(out_path)


# =====================================================
# UI
# =====================================================

class App(ttk.Frame):
    def __init__(self, root: tk.Tk):
        super().__init__(root)
        root.title("Importador de Extratos Bancários")
        root.geometry("860x420")

        self.xlsx_path = tk.StringVar()
        self.pdf_path = tk.StringVar()
        self.sheet_name = tk.StringVar()
        self.bank = tk.StringVar(value="—")

        self._build()

    def _build(self):
        self.pack(fill="both", expand=True, padx=20, pady=20)

        ttk.Label(self, text="Importador de Extratos", font=("Segoe UI", 14, "bold")).pack(anchor="w")
        ttk.Label(self, text="Selecione a planilha, aba e extrato PDF.").pack(anchor="w", pady=(0, 10))

        ttk.Button(self, text="Selecionar Planilha", command=self.pick_xlsx).pack(anchor="w")
        ttk.Entry(self, textvariable=self.xlsx_path).pack(fill="x", pady=4)

        ttk.Label(self, text="Aba").pack(anchor="w")
        self.combo = ttk.Combobox(self, textvariable=self.sheet_name, state="readonly")
        self.combo.pack(fill="x", pady=4)

        ttk.Button(self, text="Selecionar PDF", command=self.pick_pdf).pack(anchor="w")
        ttk.Entry(self, textvariable=self.pdf_path).pack(fill="x", pady=4)

        ttk.Label(self, textvariable=self.bank).pack(anchor="w", pady=4)

        ttk.Button(self, text="Processar", command=self.process).pack(anchor="e", pady=10)

    def pick_xlsx(self):
        path = filedialog.askopenfilename(filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        self.xlsx_path.set(path)
        wb = load_workbook(path, data_only=True)
        self.combo["values"] = wb.sheetnames
        self.sheet_name.set(wb.sheetnames[0])

    def pick_pdf(self):
        path = filedialog.askopenfilename(filetypes=[("PDF", "*.pdf")])
        if not path:
            return
        self.pdf_path.set(path)
        self.bank.set(f"Banco detectado: {detect_bank(path)}")

    def process(self):
        bank = detect_bank(self.pdf_path.get())

        if bank == "SICREDI":
            txs = parse_sicredi(self.pdf_path.get())
        elif bank == "PAGBANK":
            txs = parse_pagbank(self.pdf_path.get())
        elif bank == "STONE":
            txs = parse_stone(self.pdf_path.get())
        else:
            messagebox.showerror("Erro", "Banco não suportado.")
            return

        out = os.path.splitext(self.xlsx_path.get())[0] + f"_PREENCHIDO_{bank}.xlsx"
        fill_sheet(self.xlsx_path.get(), out, self.sheet_name.get(), txs, 8, self.sheet_name.get())
        messagebox.showinfo("OK", f"Arquivo gerado:\n{out}")


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()